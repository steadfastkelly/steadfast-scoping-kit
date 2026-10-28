"""Encrypt the catalog data for the site.

Builds the site's data set from the catalog, encrypts it with the team password,
and writes data/catalog.enc.js. The site decrypts it in the browser after the
password is entered. Only the encrypted file is ever published.

Two sources:

  Google Sheet (the normal path). Netlify runs this on every build:
      python3 tools/encrypt_data.py --from-sheet
  It reads three environment variables, set in Netlify's site settings:
      CATALOG_SHEET_ID             the ID in the Sheet's URL
      GOOGLE_SERVICE_ACCOUNT_JSON  the service account key (the whole JSON file)
      SITE_PASSWORD                the team password
  If any of them is missing, it keeps the committed data/catalog.enc.js and says so.

  Local workbook (fallback, or to test on a Mac):
      python3 tools/encrypt_data.py                 (reads data/catalog.xlsx, prompts for the password)
      python3 tools/encrypt_data.py --password '…'  (no prompt; single quotes if it has a $)

Changing the password: change SITE_PASSWORD in Netlify and trigger a deploy.
Everyone then needs the new one. Needs: pip3 install -r requirements.txt
"""
import json, os, sys, base64, getpass, secrets, re
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from cryptography.hazmat.primitives import hashes

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
ITER = 250_000

TABS = ["Catalog", "Quick Finder", "Packages", "Package Detail", "Level Definitions", "Terms"]

def to_dicts(table):
    """Header row plus data rows -> list of dicts, skipping blank rows."""
    if not table: return []
    hdr = [str(h) if h is not None else "" for h in table[0]]
    out = []
    for r in table[1:]:
        r = [("" if v is None else v) for v in r] + [""] * (len(hdr) - len(r))
        if any(v != "" for v in r): out.append(dict(zip(hdr, r)))
    return out

def tables_from_xlsx():
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(ROOT, "data", "catalog.xlsx"), data_only=True)
    return {n: [list(r) for r in wb[n].iter_rows(values_only=True)] for n in TABS if n in wb.sheetnames}

def tables_from_sheet(sheet_id, sa_json):
    from google.oauth2 import service_account
    from google.auth.transport.requests import AuthorizedSession
    creds = service_account.Credentials.from_service_account_info(json.loads(sa_json), scopes=["https://www.googleapis.com/auth/spreadsheets.readonly"])
    s = AuthorizedSession(creds); base = f"https://sheets.googleapis.com/v4/spreadsheets/{sheet_id}"
    have = {t["properties"]["title"] for t in s.get(base, params={"fields": "sheets.properties.title"}).json()["sheets"]}
    names = [n for n in TABS if n in have]
    r = s.get(base + "/values:batchGet", params={"ranges": [f"'{n}'" for n in names], "valueRenderOption": "UNFORMATTED_VALUE"}); r.raise_for_status()
    return {n: vr.get("values", []) for n, vr in zip(names, r.json()["valueRanges"])}

def build_data(tables):
    def rows(name): return to_dicts(tables.get(name, []))
    def lines(v): return [re.sub(r"^\s*[-•*]\s*", "", x).strip() for x in str(v or "").splitlines() if x.strip()]
    def yesno(v):
        v = str(v or "").strip().lower()
        return True if v in ("yes", "y", "true", "1") else False if v in ("no", "n", "false", "0") else None
    def num(v):
        try: return float(str(v).replace("$", "").replace(",", ""))
        except ValueError: return 0.0
    LV = {"Template": "T", "Level 1": "L1", "Level 2": "L2", "Level 3": "L3", "flat": "", "": ""}
    finder = {}
    for r in rows("Quick Finder"):
        k = r.get("Base SKU")
        if k: finder.setdefault(k, []).append(str(r.get("Client says...", "")).strip())
    REQUIRED = ["SKU", "Line item name (PandaDoc)", "Base deliverable", "Service type", "SOW section", "Level", "Level name", "Flat price (2026)"]
    hdr = [str(h).strip() for h in (tables.get("Catalog") or [[]])[0]]
    missing = [c for c in REQUIRED if c not in hdr]
    if missing:
        sys.exit("The Catalog tab is missing these column headers: " + ", ".join(missing) + "\nHeaders found: " + " | ".join(hdr))
    items = []
    for r in rows("Catalog"):
        sku = str(r["SKU"]); bs = "-".join(sku.split("-")[:3])
        items.append(dict(sku=sku, name=r["Line item name (PandaDoc)"], base=r["Base deliverable"], cat=sku.split("-")[1], category=r["Service type"],
            section=r["SOW section"], level=LV.get(str(r["Level"]), ""), lname=r["Level name"], desc=r.get("Description (customer-facing)", "") or r.get("Client description", ""),
            cdesc=str(r.get("Client description", "") or "").strip(), incl=lines(r.get("Includes")), deliv=lines(r.get("Deliverables")), note=str(r.get("Note", "") or "").strip(),
            lead=(num(r.get("Lead time (weeks)")) or None) if str(r.get("Lead time (weeks)", "")).strip() else None, mlr=yesno(r.get("MLR review")),
            price=num(r["Flat price (2026)"]), unit=r.get("Unit", ""), billing=r.get("Billing", ""), rounds=r.get("Rev rounds", ""), basis=r.get("Price basis", ""),
            use=r.get("When to use", ""), find=", ".join(finder.get(bs, [])),
            h=dict(S=num(r.get("Strategy h", 0)), D=num(r.get("Design h", 0)), V=num(r.get("Dev h", 0)), A=num(r.get("Account h", 0)), C=num(r.get("Copy h", 0)), M=num(r.get("Medical copy h", 0)))))
    cats = []; seen = set()
    for i in items:
        if i["cat"] not in seen: seen.add(i["cat"]); cats.append(dict(code=i["cat"], name=i["category"]))
    uses = {r["Package"]: r["What it is for"] for r in rows("Packages")}
    pk = {}; order = []
    for r in rows("Package Detail"):
        p, s = r.get("Package"), r.get("SKU")
        if not p or not s: continue
        if p not in pk: pk[p] = []; order.append(p)
        pk[p].append(dict(sku=str(s), qty=int(num(r.get("Qty")) or 1)))
    packs = [dict(name=p, use=uses.get(p, ""), items=pk[p]) for p in order]
    levels = []
    for r in rows("Level Definitions"):
        rounds = {m.group(1): int(m.group(2)) for m in re.finditer(r"(Template|Level \d):\s*(\d+)", str(r["Revision rounds by level"]))}
        levels.append(dict(name=r["Service type"], test=r["Level test (what makes it Level 1, 2 or 3)"], rounds=rounds))
    terms = None; rates = None; g = {}; go = []
    for r in rows("Terms"):
        if r["Group"] == "Rates": rates = r["Text"]; continue
        if r["Group"] not in g: g[r["Group"]] = []; go.append(r["Group"])
        g[r["Group"]].append([r["Heading"], r["Text"]])
    if go: terms = [dict(h=h, items=g[h]) for h in go]
    return dict(items=items, cats=cats, packs=packs, levels=levels, terms=terms, rates=rates)

def main():
    if "--from-sheet" in sys.argv:
        env = {k: os.environ.get(k, "").strip() for k in ("CATALOG_SHEET_ID", "GOOGLE_SERVICE_ACCOUNT_JSON", "SITE_PASSWORD")}
        missing = [k for k, v in env.items() if not v]
        if missing:
            print("Not set in Netlify: " + ", ".join(missing) + ". Keeping the committed data/catalog.enc.js.")
            return
        pw = env["SITE_PASSWORD"]; tables = tables_from_sheet(env["CATALOG_SHEET_ID"], env["GOOGLE_SERVICE_ACCOUNT_JSON"])
        if not tables.get("Catalog"): sys.exit("The Sheet has no Catalog tab (or it is empty). Nothing published.")
    else:
        pw = None
        if "--password" in sys.argv: pw = sys.argv[sys.argv.index("--password") + 1]
        if not pw: pw = getpass.getpass("Site password: ")
        tables = tables_from_xlsx()
    data = build_data(tables)
    plain = json.dumps(data, separators=(",", ":")).encode()
    salt = secrets.token_bytes(16); iv = secrets.token_bytes(12)
    key = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=ITER).derive(pw.encode())
    ct = AESGCM(key).encrypt(iv, plain, None)
    b64 = lambda b: base64.b64encode(b).decode()
    blob = dict(v=1, kdf="PBKDF2-SHA256", iter=ITER, salt=b64(salt), iv=b64(iv), ct=b64(ct))
    out = os.path.join(ROOT, "data", "catalog.enc.js")
    open(out, "w").write("window.SOWKIT_ENC = " + json.dumps(blob) + ";\n")
    print(f"encrypted {len(data['items'])} line items, {len(data['packs'])} packages -> data/catalog.enc.js ({os.path.getsize(out)//1024} KB)")

if __name__ == "__main__": main()
